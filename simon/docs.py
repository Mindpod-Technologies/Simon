"""Document workflow for Simon: upload → review, and create → download.

Uploads (PDF/DOCX/TXT/MD) land in ``workspace/uploads/``, their text is
extracted and ingested into the RAG store so Simon can reference the content
in any conversation. Documents Simon creates (via the ``create_document``
tool) land in ``workspace/documents/`` and are listed/downloadable in the
web UI's Documents panel.
"""

from __future__ import annotations

import logging
import re
import time
from pathlib import Path
from typing import Optional

log = logging.getLogger(__name__)

UPLOADS_DIRNAME = "uploads"
DOCUMENTS_DIRNAME = "documents"

TEXT_EXTS = {".txt", ".md", ".markdown", ".csv", ".json", ".log", ".py",
             ".yaml", ".yml", ".xml", ".html"}


def _workspace(settings) -> Path:
    base = Path(getattr(settings, "simon_workspace_dir", "./workspace"))
    base.mkdir(parents=True, exist_ok=True)
    return base


def uploads_dir(settings) -> Path:
    d = _workspace(settings) / UPLOADS_DIRNAME
    d.mkdir(parents=True, exist_ok=True)
    return d


def documents_dir(settings) -> Path:
    d = _workspace(settings) / DOCUMENTS_DIRNAME
    d.mkdir(parents=True, exist_ok=True)
    return d


def safe_name(name: str) -> str:
    """Reduce a filename to something safe to store/serve."""
    name = Path(name or "").name  # strip any directory components
    name = re.sub(r"[^\w.\- ]", "_", name).strip(". ")
    return name or f"file-{int(time.time())}"


def extract_text(path: str | Path) -> str:
    """Extract plain text from PDF/DOCX, or read text files directly."""
    path = Path(path)
    ext = path.suffix.lower()
    if ext == ".pdf":
        from pypdf import PdfReader
        reader = PdfReader(str(path))
        return "\n\n".join(
            (page.extract_text() or "") for page in reader.pages).strip()
    if ext in {".docx"}:
        import docx
        document = docx.Document(str(path))
        return "\n\n".join(
            p.text for p in document.paragraphs if p.text.strip()).strip()
    if ext in TEXT_EXTS or not ext:
        return path.read_text(encoding="utf-8", errors="replace")
    raise ValueError(f"unsupported file type: {ext or '(none)'} — "
                     f"supported: PDF, DOCX, TXT/MD and plain-text formats")


def save_upload(data: bytes, filename: str, settings,
                ingest: bool = True) -> dict:
    """Store an uploaded file, extract its text, and ingest into RAG.

    Returns ``{"name", "size", "chars", "chunks"}``. ``ingest=False`` skips
    RAG ingestion (tests). Raises ValueError on unsupported/empty content.
    """
    name = safe_name(filename)
    dest = uploads_dir(settings) / name
    dest.write_bytes(data)
    text = extract_text(dest)
    if not text.strip():
        dest.unlink(missing_ok=True)
        raise ValueError("no extractable text in that file")

    # Companion text file keeps the extracted content alongside the original.
    extracted_path = dest.with_suffix(dest.suffix + ".txt")
    extracted_path.write_text(text, encoding="utf-8")

    chunks = 0
    if ingest:
        from . import rag
        chunks = rag.add_document(extracted_path)
        log.info("uploaded %s: %d chars, %d RAG chunks", name, len(text),
                 chunks)
    return {"name": name, "size": len(data), "chars": len(text),
            "chunks": chunks}


def list_uploads(settings) -> list[dict]:
    """Uploaded originals (companion .txt files are hidden)."""
    out = []
    for p in sorted(uploads_dir(settings).iterdir()):
        if p.suffix == ".txt" and p.with_suffix("").exists():
            continue  # extraction companion
        if p.is_file():
            out.append({"name": p.name, "size": p.stat().st_size,
                        "modified": int(p.stat().st_mtime)})
    return out


def list_documents(settings) -> list[dict]:
    """Documents Simon has created, newest first."""
    docs = [{"name": p.name, "size": p.stat().st_size,
             "modified": int(p.stat().st_mtime)}
            for p in documents_dir(settings).iterdir() if p.is_file()]
    return sorted(docs, key=lambda d: -d["modified"])


def resolve_document(settings, name: str) -> Optional[Path]:
    """Resolve a document name inside documents/ — traversal-proof."""
    candidate = (documents_dir(settings) / safe_name(name)).resolve()
    if candidate.parent != documents_dir(settings).resolve():
        return None
    return candidate if candidate.is_file() else None


# ------------------------------------------------------------- create tool

def _create_document(title: str, content: str, format: str = "md",
                     settings=None) -> str:
    title = (title or "").strip()
    content = (content or "").strip()
    if not title or not content:
        return "Error: create_document needs a non-empty title and content."
    fmt = (format or "md").lower().lstrip(".")
    if fmt not in {"md", "txt", "docx"}:
        return f"Error: unsupported format '{format}' — use md, txt or docx."
    stem = safe_name(title)
    path = documents_dir(settings) / f"{stem}.{fmt}"
    if fmt == "docx":
        import docx
        document = docx.Document()
        document.add_heading(title, level=1)
        for para in content.split("\n\n"):
            if para.strip():
                document.add_paragraph(para.strip())
        document.save(str(path))
    else:
        path.write_text(
            content if fmt == "txt" else f"# {title}\n\n{content}\n",
            encoding="utf-8")
    log.info("created document %s (%d chars)", path.name, len(content))
    return (f"Created document '{path.name}' ({len(content)} chars). It is "
            f"available in the web UI Documents panel and at "
            f"workspace/documents/{path.name}.")


def register_docs_tools(registry, settings) -> None:
    """Register document-creation tools on ``registry``."""
    from .tools.base import Tool

    registry.register(Tool(
        name="create_document",
        description=(
            "Create a document (report, proposal, summary, draft) and save "
            "it for the user to download from the web UI Documents panel. "
            "Use when the user asks you to create/write/draft a document, "
            "report or file. Put the FULL content in the content parameter."),
        parameters={
            "type": "object",
            "properties": {
                "title": {"type": "string",
                          "description": "Document title (used as filename)."},
                "content": {"type": "string",
                            "description": "The complete document content."},
                "format": {"type": "string", "enum": ["md", "txt", "docx"],
                           "description": "File format; default md."},
            },
            "required": ["title", "content"],
        },
        func=lambda title, content, format="md": _create_document(
            title, content, format, settings=settings),
    ))
